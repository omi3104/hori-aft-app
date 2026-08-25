"""
ANNEX C.8 — AO Statement of Truth
Exact template match.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

import tempfile as _tempfile

def _to_bytes(doc):
    """Save document to bytes without writing to disk permanently."""
    import io as _io
    buf = _io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


NAVY = RGBColor(0x1B, 0x3A, 0x6B)


def _heading(doc, text, size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = NAVY
    return p


def _body(doc, text, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    return p


def _numbered(doc, text, size=10):
    p = doc.add_paragraph(style="List Number")
    if p.runs:
        p.runs[0].font.size = Pt(size)
        p.runs[0].text = text
    else:
        p.add_run(text).font.size = Pt(size)
    return p


def generate(fields: dict, uk_fields: dict, output_path: str, doc_date: str = ""):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    parent_name = fields.get("parent_name", "")
    parent_addr = fields.get("parent_address", "")
    uk_name     = uk_fields.get("company_name", "")
    uk_addr     = uk_fields.get("registered_address", "")
    ao_name     = fields.get("ao_full_name", "")
    ao_dob      = fields.get("ao_dob", "")
    ao_pp       = fields.get("ao_passport", "")
    ao_nat      = fields.get("ao_nationality", "")
    ao_pos      = fields.get("ao_position", "")
    date_str    = doc_date or fields.get("application_date", "")

    _heading(doc, "STATEMENT OF TRUTH")
    _heading(doc, "BY THE AUTHORISING OFFICER", size=11)
    doc.add_paragraph()

    _body(doc,
        "Reference: Sponsor Licence Application – UK Expansion Worker Route (Global Business Mobility)",
        size=10, bold=True)
    doc.add_paragraph()

    _body(doc, f"Business Name:", size=10, bold=True)
    _body(doc, f"{parent_name}  Address: {parent_addr}", size=10)
    doc.add_paragraph()
    _body(doc, f"{uk_name} (Subsidiary Registered Office)  Address: {uk_addr}", size=10)
    doc.add_paragraph()

    _body(doc,
        f"Applicant (Authorising Officer): {ao_name}, D.O.B: {ao_dob}, "
        f"Passport Number {ao_pp}, {ao_nat}.", size=10)
    doc.add_paragraph()

    _body(doc,
        f"I, {ao_name}, hereby make the following statement in connection with the "
        f"above-mentioned Sponsor Licence application and do so truthfully and on behalf of "
        f"{parent_name}.", size=10)
    doc.add_paragraph()

    statements = [
        f"I am currently employed by {parent_name} and have been continuously employed by the "
        f"company for a period exceeding twelve (12) months prior to the date of this statement.",

        f"I have been formally appointed as the Authorising Officer for the proposed establishment "
        f"of the UK subsidiary of {parent_name}. This subsidiary has been incorporated under the "
        f"laws of England and Wales in accordance with section 1159 of the Companies Act 2006, as "
        f"a wholly owned subsidiary of our overseas parent company based in Pakistan.",

        f"{parent_name} has engaged the services of Chisty Law Chambers LLP (Incorporation No. "
        f"0269333), a legal consultancy firm based in Pakistan, for the purposes of providing "
        f"general advisory support in relation to the Sponsor Licence application. Their principal "
        f"business place is located at: 2nd Floor, Almas Tower, MM Alam Road, Gulberg II, Lahore, Pakistan.",

        f"The scope of services provided by Chisty Law Chambers LLP is strictly limited to general "
        f"advisory, compilation, and procedural assistance in connection with the preparation of the "
        f"company's business documentation. This includes, without limitation, the Business Plan, "
        f"twelve-month financial projections, and the Sponsorship Licence application. For the "
        f"avoidance of doubt, the firm 'Chisty Law Chambers LLP' does not assume responsibility "
        f"for the accuracy, completeness, or substantive content of any such documentation, which "
        f"remains the sole responsibility of the authorising officer (me) and/or the company.",

        f"Chisty Law Chambers LLP has not been appointed to act as our legal representative for "
        f"the purposes of this application, nor are they authorised to submit the application to "
        f"the UKVI on my behalf or on behalf of the company.",

        f"I hereby acknowledge and confirm that I, and the company which I represent, bear full "
        f"and continuing responsibility for the accuracy, completeness, and truthfulness of all "
        f"information and documentation submitted in connection with the UK Expansion Worker Visa "
        f"Sponsorship Licence Application.",

        f"This responsibility expressly includes, without limitation, all business and personal "
        f"documentation supplied by or on behalf of the company or myself, including (but not "
        f"limited to) the Business Plan, 12-month financial projections, and any and all supporting "
        f"materials provided in support of the application.",

        f"I further acknowledge that such information and documentation have been prepared and "
        f"provided under my instruction and authority.",

        f"I further confirm that the contents of this statement are true to the best of my "
        f"knowledge and belief.",
    ]

    for s in statements:
        _numbered(doc, s)
        doc.add_paragraph()

    _body(doc, "Thank You,", size=10)
    doc.add_paragraph()
    _body(doc, "Signed: ___________________________", size=10)
    doc.add_paragraph()
    _body(doc, f"{ao_name} – Authorising Officer", size=10, bold=True)
    doc.add_paragraph()
    _body(doc, "For and on behalf of", size=10)
    doc.add_paragraph()
    _body(doc, f"{parent_name} (Pakistan Parent Company) &", size=10)
    doc.add_paragraph()
    _body(doc, f"{uk_name} (UK Subsidiary)", size=10)
    doc.add_paragraph()
    _body(doc, f"Date: {date_str}", size=10)

    if output_path is None:
        return _to_bytes(doc)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
