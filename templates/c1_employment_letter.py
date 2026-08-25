"""
ANNEX C.1 — Letter from Company HR, confirming Employment
Issued by: Pakistan Parent Company
Font: Arial 11pt, space_after=10pt (matches real document)
"""
import io, os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = "Arial"

def _to_bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()

def _set_doc_font(doc):
    doc.styles['Normal'].font.name = FONT

def _p(doc, text="", bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
    return p

def generate(fields: dict, uk_fields: dict, output_path,
             salary: str = "", start_date: str = "", doc_date: str = ""):
    doc = Document()
    _set_doc_font(doc)
    for sec in doc.sections:
        sec.top_margin    = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin   = Cm(3)
        sec.right_margin  = Cm(2.5)

    parent_name = fields.get("parent_name", "")
    parent_addr = fields.get("parent_address", "")
    ao_name     = fields.get("ao_full_name", "")
    ao_pp       = fields.get("ao_passport", "")
    ao_pos      = fields.get("ao_position", "")
    ao_uk_title = fields.get("ao_uk_title", "Executive Director")
    directors   = fields.get("parent_directors", [])
    job_desc    = fields.get("job_description", "")
    job_duties  = fields.get("job_duties", [])
    department  = fields.get("department", "")
    app_date    = doc_date or fields.get("application_date", "")

    # Signatory — use a director who is not the AO
    signatory = next((d for d in directors if d.strip().lower() != ao_name.strip().lower()), "")
    if not signatory and directors:
        signatory = directors[-1]

    # ── Letterhead ──
    _p(doc, "Confidential", bold=True, size=11)
    _p(doc, parent_name, bold=True, size=11)
    if parent_addr:
        for line in parent_addr.split(","):
            l = line.strip()
            if l:
                _p(doc, l, size=11, space_after=0)
    _p(doc)
    _p(doc, f"Date: {app_date}" if app_date else "Date: ___ / ___ / 20___", size=11)
    _p(doc)

    _p(doc, "To Whom It May Concern,", bold=True, size=11)
    _p(doc)

    _p(doc, f"Re: Employment Confirmation Letter – Mr. {ao_name}", bold=True, size=11)
    _p(doc)

    # Paragraph 1 – confirm employment
    p1 = doc.add_paragraph()
    p1.paragraph_format.space_after = Pt(10)
    p1.paragraph_format.space_before = Pt(0)
    r = p1.add_run(
        f"This is to certify that Mr. {ao_name}, holder of Passport No. {ao_pp}"
        + (f", has been employed with {parent_name} since {start_date}"
           if start_date else
           f", has been associated with {parent_name} since its inception")
        + " and continues to be in active employment with the business."
    )
    r.font.name = FONT
    r.font.size = Pt(11)

    _p(doc)

    # Paragraph 2 – role and company context
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(10)
    p2.paragraph_format.space_before = Pt(0)
    role_desc = f"the {ao_pos}" if ao_pos else "a senior executive"
    r = p2.add_run(
        f"Mr. {ao_name} currently serves as {role_desc} of the business. "
        f"The business is duly registered and operates in compliance with all applicable "
        f"regulatory requirements."
    )
    r.font.name = FONT
    r.font.size = Pt(11)

    _p(doc)

    # Paragraph 3 – responsibilities
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(10)
    p3.paragraph_format.space_before = Pt(0)
    if job_desc:
        r = p3.add_run(
            f"In his capacity as {ao_pos or ao_uk_title}, Mr. {ao_name} is responsible for "
            f"{job_desc}"
        )
    else:
        r = p3.add_run(
            f"In his capacity as {ao_pos or ao_uk_title}, Mr. {ao_name} is responsible for "
            f"providing overall strategic leadership and direction to the business. He oversees "
            f"all business operations, client relationship management, financial performance, "
            f"regulatory compliance, and ensuring transparent and timely delivery of services."
        )
    r.font.name = FONT
    r.font.size = Pt(11)

    if job_duties:
        _p(doc, "His key responsibilities include:", size=11)
        for duty in job_duties:
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_after = Pt(2)
            r = bp.add_run(duty)
            r.font.name = FONT
            r.font.size = Pt(11)

    _p(doc)

    # Paragraph 4 – character
    p4 = doc.add_paragraph()
    p4.paragraph_format.space_after = Pt(10)
    p4.paragraph_format.space_before = Pt(0)
    r = p4.add_run(
        f"Mr. {ao_name} is a dedicated and principled professional who consistently works "
        f"to the highest standards of integrity and professionalism."
    )
    r.font.name = FONT
    r.font.size = Pt(11)

    _p(doc)
    _p(doc,
       f"This letter is issued at the request of Mr. {ao_name} for official purposes. "
       f"The contents of this letter are verifiable and may be confirmed by contacting "
       f"our HR department directly.",
       size=11)
    _p(doc)
    _p(doc, "Thank You,", size=11)
    _p(doc)
    _p(doc)
    _p(doc, "Signed: ___________________________", size=11)
    _p(doc)
    if signatory:
        _p(doc, f"{signatory} – HR & Admin Manager", bold=True, size=11)
    _p(doc, "For and on behalf of", size=11)
    _p(doc, parent_name, bold=True, size=11)

    if output_path is None:
        return _to_bytes(doc)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
