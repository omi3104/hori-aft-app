"""
ANNEX C.2 — Board Resolution / Minutes of Meeting from the Parent Company
Font: Arial 11pt, 1.5x line spacing (matches real document)
"""
import io, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Arial"

def _to_bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()

def _set_doc_font(doc):
    doc.styles['Normal'].font.name = FONT

def _p(doc, text="", bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
    return p

def _heading(doc, text, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = True
    return p

def generate(fields: dict, uk_fields: dict, output_path,
             meeting_date: str = "", directors: list = None, attendees: list = None):
    doc = Document()
    _set_doc_font(doc)
    for sec in doc.sections:
        sec.top_margin    = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin   = Cm(3)
        sec.right_margin  = Cm(2.5)

    parent_name  = fields.get("parent_name", "")
    parent_reg   = fields.get("parent_reg_no", "")
    parent_ref   = fields.get("parent_ref_no", "")
    parent_addr  = fields.get("parent_address", "")
    ao_name      = fields.get("ao_full_name", "")
    ao_pos       = fields.get("ao_position", "")
    ao_uk_title  = fields.get("ao_uk_title", "Executive Director")
    uk_name      = uk_fields.get("company_name", "")
    uk_inc_date  = uk_fields.get("incorporation_date", "")
    directors    = directors or fields.get("parent_directors", [])
    mtg_date     = meeting_date or fields.get("application_date", "")

    # Use attendees if provided (list of "Name, Designation" strings), else fall back to directors
    # attendees takes priority — extracted from hierarchy chart and confirmed in review
    attendance_list = attendees if attendees else directors

    # Chairperson = first attendee (or first director, or AO)
    chairperson_line = attendance_list[0] if attendance_list else ao_name
    # Extract just the name part (before the comma) for "chaired by" line
    chairperson_name = chairperson_line.split(",")[0].strip() if "," in chairperson_line else chairperson_line.strip()

    # ── Header ──
    _p(doc, parent_name, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    if parent_reg:
        _p(doc, f"Company Registration No.: {parent_reg}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    if parent_ref:
        _p(doc, f"Company Reference No.: {parent_ref}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    if parent_addr:
        _p(doc, f"Registered Address: {parent_addr}", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc)

    _p(doc, "BOARD RESOLUTION / MINUTES OF THE MEETING", bold=True, size=12,
       align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    # ── Date / Time / Location / Agenda table (matches real doc) ──
    mtbl = doc.add_table(rows=0, cols=2)
    mtbl.style = "Table Grid"
    mtbl.columns[0].width = Cm(4)
    mtbl.columns[1].width = Cm(12)
    for lbl, val in [
        ("Date",     mtg_date),
        ("Time",     "10:00 AM"),
        ("Location", parent_addr),
        ("Agenda",   "Appointment of Authorising Officer for UK Business Establishment and "
                     "Immigration Sponsorship Licence Application."),
    ]:
        row = mtbl.add_row()
        for cell, txt, bold in [(row.cells[0], lbl, True), (row.cells[1], val, False)]:
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.5
            r = p.add_run(txt)
            r.font.name = FONT
            r.font.size = Pt(11)
            r.bold = bold

    _p(doc)

    # 1. Attendance
    _heading(doc, "1. Attendance")
    _p(doc, "The following company officials were present:", size=11)

    if attendance_list:
        for person in attendance_list:
            bp = doc.add_paragraph()
            bp.paragraph_format.space_after = Pt(2)
            bp.paragraph_format.line_spacing = 1.5
            r = bp.add_run(person)
            r.font.name = FONT
            r.font.size = Pt(11)
    else:
        _p(doc, "[Attendees to be inserted]", size=11)
    _p(doc)

    # 2. Chairperson
    _heading(doc, "2. Chairperson")
    _p(doc, f"The meeting was chaired by {chairperson_name}.", size=11)
    _p(doc)

    # 3. Purpose
    _heading(doc, "3. Purpose of the Meeting")
    _p(doc,
       "To consider and approve the appointment of an Authorising Officer for the purposes of "
       "establishing the company in the United Kingdom and making an application to the Home Office "
       "for a UK Sponsor Licence under the applicable UK Immigration Rules.", size=11)
    _p(doc)

    # 4. Discussion
    _heading(doc, "4. Discussion")
    _p(doc,
       "The Chairperson explained the importance of appointing an Authorising Officer, as required "
       "by UK Visas and Immigration (UKVI), to ensure compliance with all sponsorship obligations, including:", size=11)

    bullets = [
        "Managing the Sponsor Licence Application Process.",
        "Acting as the Main Point of Contact with The Home Office.",
        "Overseeing Immigration Compliance and Record-Keeping Requirements.",
        "Being A Senior and Responsible Person Based Permanently in the UK.",
    ]
    for b in bullets:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(2)
        bp.paragraph_format.line_spacing = 1.5
        r = bp.add_run(b)
        r.font.name = FONT
        r.font.size = Pt(11)

    _p(doc)
    _p(doc,
       "The board acknowledged that the role carries legal responsibilities and must be held by a "
       "senior, suitable and trusted individual within the company.", size=11)
    _p(doc)

    # 5. Resolution
    _heading(doc, "5. Resolution")
    _p(doc,
       "After due consideration, based on the education, professional experience and competence, "
       "the board unanimously resolved to appoint:", size=11)
    _p(doc)
    _p(doc, f"Name:      {ao_name}", bold=True, size=11)
    _p(doc, f"Position:  {ao_pos or ao_uk_title}", bold=True, size=11)
    _p(doc)
    _p(doc,
       f"As the Authorising Officer of the UK subsidiary office with name \"{uk_name}\", "
       f"which has been registered already on {uk_inc_date}, for the purposes outlined above, "
       f"with immediate effect.", size=11)
    _p(doc)
    _p(doc,
       "The appointed Authorising Officer is authorised and directed to take immediate and all "
       "requisite steps to commence and complete the Sponsor Licence application process.", size=11)
    _p(doc)
    _p(doc,
       "The Authorising Officer shall also be responsible for ensuring that the company fully "
       "complies with all ongoing duties, responsibilities, and obligations set out in the UK "
       "Visas and Immigration (UKVI) Sponsor Guidance.", size=11)
    _p(doc)

    # 6. Closing
    _heading(doc, "6. Closing")
    _p(doc, "There being no further business, the meeting was concluded at 12:30 PM.", size=11)
    _p(doc)
    _p(doc, "Approved by:", bold=True, size=11)
    _p(doc)
    _p(doc)
    _p(doc, "Signed: ___________________________", size=11)
    _p(doc, f"{chairperson_name} – Chairperson", size=11)
    _p(doc)
    _p(doc, "For and on behalf of", size=11)
    _p(doc, parent_name, bold=True, size=11)

    if output_path is None:
        return _to_bytes(doc)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
