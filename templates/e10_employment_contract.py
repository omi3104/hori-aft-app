"""
ANNEX E.10 — Draft Employment Contract
Exact template match.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

import tempfile as _tempfile

def _to_bytes(doc):
    """Save document to bytes without writing to disk permanently."""
    import io as _io
    buf = _io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


NAVY = RGBColor(0x1B, 0x3A, 0x6B)


def _heading(doc, text, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = NAVY
    return p


def _body(doc, text, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
    return p


def _alpha(doc, text, size=10):
    p = doc.add_paragraph(style="List Continue")
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def _bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def generate(fields: dict, uk_fields: dict, output_path: str,
             salary: str = "", start_date: str = ""):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    uk_name   = uk_fields.get("company_name", "")
    uk_num    = uk_fields.get("company_number", "")
    uk_addr   = uk_fields.get("registered_address", "")
    parent_name = fields.get("parent_name", "")
    ao_name   = fields.get("ao_full_name", "")
    ao_dob    = fields.get("ao_dob", "")
    ao_pp     = fields.get("ao_passport", "")
    ao_nat    = fields.get("ao_nationality", "")
    ao_uk_title = fields.get("ao_uk_title", "Executive Director")
    soc_code  = fields.get("ao_soc_code", "1111")
    sal       = salary or fields.get("ao_going_rate", "£60,000")
    sal_hr    = fields.get("ao_going_rate_hourly", "£30.77 per hour")

    _heading(doc, "Employment Agreement", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()

    _body(doc, "This Employment Agreement is made on (Insert Date) between:", size=10)
    doc.add_paragraph()

    _body(doc,
        f"Employer: {uk_name}: Registered Office: {uk_addr}  "
        f"Company Number: {uk_num} (the \"Employer\" or \"Company\") "
        f"(Registered Subsidiary of {parent_name}, a Proprietorship Business "
        f"registered in Pakistan)", size=10, bold=False)
    doc.add_paragraph()

    _body(doc, "And", size=10)
    doc.add_paragraph()

    _body(doc,
        f"Employee: {ao_name}, D.O.B: {ao_dob}, Passport Number {ao_pp}, "
        f"{ao_nat}. (the \"Employee\")", size=10)
    doc.add_paragraph()

    # Clauses
    _heading(doc, "1. Commencement and Duration")
    _alpha(doc,
        "The Employee's employment under this Agreement shall commence on (Insert Date) and "
        "shall continue until one year from Commencement Date, unless terminated earlier in "
        "accordance with the provisions of this Agreement.")
    _alpha(doc,
        "The Employee's employment is subject to the Employee obtaining and maintaining a valid "
        "UK Expansion Worker visa under the Global Business Mobility: Expansion Worker immigration "
        "route, sponsored by the Employer. The Employee agrees to provide all necessary "
        "documentation and cooperate fully in the visa application and maintenance process.")
    _alpha(doc,
        "This Agreement constitutes a fixed-term contract for the purposes of the Employer's "
        "sponsorship obligations under the UK Immigration Rules.")

    doc.add_paragraph()
    _heading(doc, "2. Job Title & Duties")
    _alpha(doc,
        f"The Employee shall be employed as {ao_uk_title}, which corresponds to the Standard "
        f"Occupational Classification (SOC) 2020 occupation code {soc_code} as specified on the "
        f"Certificate of Sponsorship (CoS) issued by the Employer.")
    _alpha(doc,
        "The Employee's principal duties shall be to establish and manage the UK branch "
        "operations, including market entry strategies, sales, and compliance with the proposed "
        "business plan in accordance with UK regulations.")

    doc.add_paragraph()
    _heading(doc, "3. Place of Work")
    _alpha(doc,
        f"The Employee's principal place of work shall be {uk_addr}, United Kingdom and/or such "
        f"other location in the UK as the Employer may reasonably require.")
    _alpha(doc,
        "The Employee may be required to travel within the UK or overseas as necessary for the "
        "performance of their duties, with reasonable expenses reimbursed in accordance with the "
        "Employer's policies.")

    doc.add_paragraph()
    _heading(doc, "4. Remuneration")
    _alpha(doc,
        f"The Employee shall be paid a gross annual salary of {sal} payable in monthly instalments "
        f"in arrears by bank transfer or any other banking transaction modes on or before 5th "
        f"working day of every month.")
    _alpha(doc, "The Salary is inclusive of all allowances unless specified otherwise below.")
    _alpha(doc,
        "The Salary shall be subject to deductions for tax, National Insurance contributions, "
        "and any other statutory deductions. The Employee will be paid through the Employer's "
        "PAYE scheme.")
    _alpha(doc,
        "The Employer confirms that the remuneration complies with the National Minimum Wage "
        "Act 1998 and meets or exceeds the minimum salary threshold for the UK Expansion Worker route.")

    doc.add_paragraph()
    _heading(doc, "5. Hours of Work")
    _alpha(doc,
        "The Employee's normal working hours shall be 37.5 hours per week, Monday to Friday, "
        "9:00 am to 5:30 pm, exclusive of lunch breaks.")

    doc.add_paragraph()
    _heading(doc, "6. Holiday Entitlement")
    _alpha(doc,
        "The Employee is entitled to 28 days paid annual leave per holiday year, including "
        "the public holidays.")
    _alpha(doc,
        "Holiday must be taken at times approved by the Employer and may not be carried "
        "forward without prior agreement.")

    doc.add_paragraph()
    _heading(doc, "7. Pension and Benefits")
    _alpha(doc,
        "The Employer shall enroll the Employee in its auto-enrolment pension scheme in "
        "accordance with the Pensions Act 2008.")

    doc.add_paragraph()
    _heading(doc, "8. Sponsorship and Immigration Obligations")
    _alpha(doc,
        "The Employee acknowledges that their employment is conditional upon the Employer "
        "holding a valid Sponsor Licence and issuing a valid CoS for the UK Expansion Worker route.")
    _alpha(doc, "The Employee shall:")
    for item in [
        "Work only for the Employer in the role specified on the CoS.",
        "Not engage in supplementary employment except as permitted under the Immigration Rules.",
        "Not access public funds and shall maintain sufficient funds for themselves and any dependents.",
        "Promptly notify the Employer of any changes in personal circumstances that may affect their visa status.",
        "Apply for visa extensions or changes as directed by the Employer, with associated costs borne by the Employer.",
    ]:
        _bullet(doc, item)

    _alpha(doc,
        "If the Employee's visa or sponsorship is revoked or expires without renewal, the "
        "employment shall terminate automatically.")
    _alpha(doc,
        "The Employer may terminate this Agreement immediately without notice if the Employee "
        "breaches immigration conditions or if the Sponsor's licence is downgraded, curtailed, "
        "or revoked.")

    doc.add_paragraph()
    _heading(doc, "9. Termination")
    _alpha(doc,
        "Either party may terminate this Agreement by giving one month's written notice.")
    _alpha(doc,
        "The Employer may terminate this Agreement immediately (without notice) in cases of "
        "gross misconduct, serious breach of contract, or breach of immigration conditions.")

    doc.add_paragraph()
    _heading(doc, "10. Governing Law")
    _alpha(doc,
        "This Agreement shall be governed by and construed in accordance with the laws of "
        "England and Wales. Any disputes shall be subject to the exclusive jurisdiction of "
        "the courts of England and Wales.")

    doc.add_paragraph()
    _body(doc, "Signed by the Employer:", size=10, bold=True)
    _body(doc, "Signed: ___________________________", size=10)
    _body(doc, f"{ao_name} – {ao_uk_title}", size=10)
    _body(doc, f"For and on behalf of {uk_name}", size=10)
    doc.add_paragraph()
    _body(doc, "Signed by the Employee:", size=10, bold=True)
    _body(doc, "Signed: ___________________________", size=10)
    _body(doc, f"{ao_name}", size=10)
    doc.add_paragraph()
    _body(doc, "Date: (Insert Date)", size=10)

    if output_path is None:
        return _to_bytes(doc)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
