"""
ANNEX A.1 — Title Pages
Exact template: navy header table, client/UK subsidiary details, annex index.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

import tempfile as _tempfile

def _to_bytes(doc):
    """Save document to bytes without writing to disk permanently."""
    import io as _io
    buf = _io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


NAVY = RGBColor(0x1B, 0x3A, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)


def _set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _cell_para(cell, text, bold=False, size=11, color=WHITE, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.paragraphs[0].clear()
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p


def _add_navy_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = NAVY
    return p


def _add_detail_row(table, label, value):
    row = table.add_row()
    label_cell = row.cells[0]
    value_cell = row.cells[1]
    _set_cell_bg(label_cell, "1B3A6B")
    lp = label_cell.paragraphs[0]
    lp.clear()
    lr = lp.add_run(label)
    lr.bold = True
    lr.font.size = Pt(10)
    lr.font.color.rgb = WHITE
    vp = value_cell.paragraphs[0]
    vp.clear()
    vr = vp.add_run(value)
    vr.font.size = Pt(10)
    vr.font.color.rgb = BLACK


def generate(fields: dict, uk_fields: dict, output_path: str):
    """
    fields: parent company data
    uk_fields: UK subsidiary data (from Companies House)
    output_path: full path to save the .docx
    """
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ── PAGE 1: Main Title Page ──────────────────────────────────────────────
    _add_navy_heading(doc, "ANNEX A")
    _add_navy_heading(doc, "SPONSOR LICENCE APPLICATION")
    _add_navy_heading(doc, "UK EXPANSION WORKER ROUTE")
    _add_navy_heading(doc, "(GLOBAL BUSINESS MOBILITY)")
    doc.add_paragraph()

    # Parent company table
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.columns[0].width = Inches(2.2)
    t.columns[1].width = Inches(4.3)
    hdr = t.rows[0]
    hdr.cells[0].merge(hdr.cells[1])
    _set_cell_bg(hdr.cells[0], "1B3A6B")
    _cell_para(hdr.cells[0], "PARENT COMPANY", bold=True, size=11,
               align=WD_ALIGN_PARAGRAPH.CENTER)

    parent_rows = [
        ("Business / Company Name", fields.get("parent_name", "")),
        ("Registration Number",     fields.get("parent_reg_no", "")),
        ("Reference No",            fields.get("parent_ref_no", "")),
        ("Registered On",           fields.get("parent_reg_date", "")),
        ("Registered Address",      fields.get("parent_address", "")),
    ]
    for label, val in parent_rows:
        _add_detail_row(t, label, val)

    doc.add_paragraph()

    # UK subsidiary table
    t2 = doc.add_table(rows=1, cols=2)
    t2.style = "Table Grid"
    t2.columns[0].width = Inches(2.2)
    t2.columns[1].width = Inches(4.3)
    hdr2 = t2.rows[0]
    hdr2.cells[0].merge(hdr2.cells[1])
    _set_cell_bg(hdr2.cells[0], "1B3A6B")
    _cell_para(hdr2.cells[0], "UK SUBSIDIARY", bold=True, size=11,
               align=WD_ALIGN_PARAGRAPH.CENTER)

    uk_rows = [
        ("Business / Company Name", uk_fields.get("company_name", "")),
        ("Company Number",          uk_fields.get("company_number", "")),
        ("Incorporated On",         uk_fields.get("incorporation_date", "")),
        ("Registered Address",      uk_fields.get("registered_address", "")),
        ("SIC Codes",               ", ".join(uk_fields.get("sic_codes", []))),
    ]
    for label, val in uk_rows:
        _add_detail_row(t2, label, val)

    doc.add_paragraph()

    # AO details
    ao_name = fields.get("ao_full_name", "")
    ao_dob  = fields.get("ao_dob", "")
    ao_pp   = fields.get("ao_passport", "")
    ao_nat  = fields.get("ao_nationality", "")

    t3 = doc.add_table(rows=1, cols=2)
    t3.style = "Table Grid"
    t3.columns[0].width = Inches(2.2)
    t3.columns[1].width = Inches(4.3)
    hdr3 = t3.rows[0]
    hdr3.cells[0].merge(hdr3.cells[1])
    _set_cell_bg(hdr3.cells[0], "1B3A6B")
    _cell_para(hdr3.cells[0], "AUTHORISING OFFICER", bold=True, size=11,
               align=WD_ALIGN_PARAGRAPH.CENTER)

    ao_rows = [
        ("Full Name",       ao_name),
        ("Date of Birth",   ao_dob),
        ("Passport Number", ao_pp),
        ("Nationality",     ao_nat),
        ("Job Title",       fields.get("ao_position", "")),
    ]
    for label, val in ao_rows:
        _add_detail_row(t3, label, val)

    # Page break → Annex Index Page
    doc.add_page_break()

    _add_navy_heading(doc, "ANNEX INDEX")
    doc.add_paragraph()

    annexes = [
        ("ANNEX A", "General / Application Information"),
        ("ANNEX B", "Consultancy Agreement"),
        ("ANNEX C", "Authorising Officer Documents"),
        ("ANNEX D", "Business Plan & Financial Projections"),
        ("ANNEX E", "Employment Contract"),
        ("ANNEX F", "Supporting Documents"),
    ]

    t4 = doc.add_table(rows=1, cols=2)
    t4.style = "Table Grid"
    t4.columns[0].width = Inches(1.5)
    t4.columns[1].width = Inches(5.0)
    hdr4 = t4.rows[0]
    _set_cell_bg(hdr4.cells[0], "1B3A6B")
    _set_cell_bg(hdr4.cells[1], "1B3A6B")
    _cell_para(hdr4.cells[0], "ANNEX", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell_para(hdr4.cells[1], "DESCRIPTION", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    for code, desc in annexes:
        row = t4.add_row()
        rp = row.cells[0].paragraphs[0]
        rp.clear()
        rr = rp.add_run(code)
        rr.bold = True
        rr.font.size = Pt(10)
        rr.font.color.rgb = NAVY
        dp = row.cells[1].paragraphs[0]
        dp.clear()
        dr = dp.add_run(desc)
        dr.font.size = Pt(10)

    doc.add_paragraph()
    doc.add_paragraph()

    # Closing signature block
    p = doc.add_paragraph()
    p.add_run("Signed: ___________________________").bold = False
    doc.add_paragraph()
    closing = (
        f"For and on behalf of\n\n"
        f"{fields.get('parent_name', '')} (Pakistan Parent Company) &\n\n"
        f"{uk_fields.get('company_name', '')} (UK Subsidiary)"
    )
    for line in closing.split("\n"):
        lp = doc.add_paragraph(line)
        lp.runs[0].font.size = Pt(10) if lp.runs else None

    if output_path is None:
        return _to_bytes(doc)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
