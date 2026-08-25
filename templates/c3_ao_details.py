"""
ANNEX C.3 — AO Details Letter
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

import tempfile as _tempfile

def _to_bytes(doc):
    """Save document to bytes without writing to disk permanently."""
    import io as _io
    buf = _io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


NAVY = RGBColor(0x1B, 0x3A, 0x6B)


def _body(doc, text, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    return p


def generate(fields: dict, uk_fields: dict, output_path: str):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    parent_name = fields.get("parent_name", "")
    uk_name     = uk_fields.get("company_name", "")
    ao_name     = fields.get("ao_full_name", "")
    ao_dob      = fields.get("ao_dob", "")
    ao_pp       = fields.get("ao_passport", "")
    ao_nat      = fields.get("ao_nationality", "")
    ao_pos      = fields.get("ao_position", "")
    ao_uk_title = fields.get("ao_uk_title", "Executive Director")
    soc_code    = fields.get("ao_soc_code", "1111")
    going_rate  = fields.get("ao_going_rate", "£60,000")
    going_hourly= fields.get("ao_going_rate_hourly", "£30.77 per hour")

    # Heading
    _body(doc, "To:", size=10)
    _body(doc, "Sponsor Casework Operations", size=10)
    _body(doc, "Vulcan House - Steel", size=10)
    _body(doc, "PO Box 3468", size=10)
    _body(doc, "Sheffield S3 8WA, United Kingdom", size=10)
    doc.add_paragraph()
    _body(doc, "Dear Sirs:", size=10)
    doc.add_paragraph()

    _body(doc,
        "Ref: Details of Identified Authorising Officer for Subsidiary Office – "
        "UK Expansion Worker Route (Global Business Mobility)", size=10, bold=True)
    doc.add_paragraph()

    _body(doc,
        f"Applicant (Authorising Officer): {ao_name}, D.O.B: {ao_dob}, "
        f"Passport Number {ao_pp}, {ao_nat}.", size=10)
    doc.add_paragraph()

    _body(doc,
        "We write in reference to the above. Please see below the details of the required "
        "job role for our Authorising Officer.", size=10)
    doc.add_paragraph()

    _body(doc,
        "The details are being provided in accordance with the requirements mentioned in "
        "paragraph 3.10 of Appendix A-Supporting Documents for Sponsor Licence.", size=10)
    doc.add_paragraph()

    _body(doc, "We are requesting the following job role for our Authorising Officer:", size=10)
    doc.add_paragraph()

    job_details = [
        ("Job Title",         ao_uk_title),
        ("Job Type",          "Chief Executive and other senior officials"),
        ("SOC Code",          soc_code),
        ("Going Rate",        f"{going_rate} ({going_hourly})"),
    ]
    for label, val in job_details:
        _body(doc, f"{label}: {val}", size=10)

    doc.add_paragraph()
    _body(doc, "The Authorising Officer/Worker has been identified considering:", size=10)
    criteria = [
        "The seniority and important role in the overseas business",
        "The area of operations in the overseas business",
        "The knowledge of relevant industry and services",
        "The educational background, certifications and communication skills",
        "The proposed operations of the UK entity and relevance of the worker",
    ]
    for c in criteria:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(c).font.size = Pt(10)

    doc.add_paragraph()
    _body(doc,
        f"Moreover, Subject to grant of sponsor license and immigration permission to "
        f"{ao_name}, he will travel to the United Kingdom to undertake the assignment of expansion.",
        size=10)
    doc.add_paragraph()

    _body(doc,
        "We believe that our Sponsor Licence Application fulfils all the requirements and is "
        "presented in a manner consistent with the standards and expectations relevant to "
        "establishing a subsidiary office in the United Kingdom.", size=10)
    doc.add_paragraph()

    _body(doc,
        "However, should you have any further queries or require clarification or additional "
        "supporting documentation in relation to our application, please do not hesitate to "
        "contact us.", size=10)
    doc.add_paragraph()

    _body(doc, "Thank You,", size=10)
    doc.add_paragraph()
    _body(doc, "Signed: ___________________________", size=10)
    doc.add_paragraph()

    directors = fields.get("parent_directors", [])
    signatory = next((d for d in directors if d != ao_name), ao_name)
    _body(doc, f"{signatory}", size=10, bold=True)
    doc.add_paragraph()
    _body(doc, "For and on behalf of", size=10)
    doc.add_paragraph()
    _body(doc, f"{parent_name} (Pakistan Parent Company) &", size=10)
    doc.add_paragraph()
    _body(doc, f"{uk_name} (UK Subsidiary)", size=10)

    if output_path is None:
        return _to_bytes(doc)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
